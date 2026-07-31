"""분류 결과 내보내기 유틸 (Excel 생성, 대표도면 삽입).

streamlit에 의존하지 않아 단독으로 테스트할 수 있다.
"""
from __future__ import annotations

import json
from pathlib import Path


def _drawing_paths(drawings_json: str | None, limit: int = 12) -> list[str]:
    """도면 목록에서 실제 존재하는 파일 경로를 순서대로 반환 (대표도면 우선)."""
    if not drawings_json:
        return []
    try:
        drawings = json.loads(drawings_json)
    except (json.JSONDecodeError, TypeError):
        return []
    if not isinstance(drawings, list):
        return []

    ordered = [d for d in drawings if isinstance(d, dict) and d.get("drawing_type") == "representative"]
    ordered += [d for d in drawings if isinstance(d, dict) and d not in ordered]

    paths = []
    for d in ordered:
        p = d.get("file_path") or d.get("url")
        if p and Path(p).exists():
            paths.append(p)
        if len(paths) >= limit:
            break
    return paths


def _file_hash(path: str, _cache: dict = {}) -> str | None:
    """파일 내용 해시 (동일 이미지 판별용). 경로별로 캐시."""
    import hashlib

    if path in _cache:
        return _cache[path]
    try:
        h = hashlib.md5(Path(path).read_bytes()).hexdigest()
    except OSError:
        h = None
    _cache[path] = h
    return h


def _color_stats(path: str, _cache: dict = {}) -> tuple[float, float, int]:
    """(평균 채도, 유채색 픽셀 비율, 픽셀 수). 장식 이미지 판별용."""
    if path in _cache:
        return _cache[path]
    try:
        from PIL import Image as PILImage

        with PILImage.open(path) as im:
            im = im.convert("RGB")
            im.thumbnail((80, 80))  # 통계용이므로 축소해 빠르게
            w, h = im.size
            hist = im.convert("HSV").split()[1].histogram()  # 채도 채널 분포
        total = sum(hist) or 1
        mean_sat = sum(i * c for i, c in enumerate(hist)) / total
        colorful = sum(c for i, c in enumerate(hist) if i > 40) / total
        result = (mean_sat, colorful, w * h)
    except Exception:
        result = (0.0, 0.0, 0)
    _cache[path] = result
    return result


def _image_score(path: str, repeat_count: int) -> float:
    """대표도면 적합도 점수. 높을수록 실제 도면일 가능성이 큼.

    - 여러 건에 반복 등장할수록 감점 (서식·표지 이미지)
    - 사진처럼 채도가 높으면 감점 (EU 깃발 등 장식 배경)
    - 지나치게 작으면 감점 (로고·아이콘)
    """
    score = 100.0
    if repeat_count >= 2:
        score -= 25.0 * min(repeat_count, 6)

    mean_sat, colorful, area = _color_stats(path)
    if mean_sat > 60:          # 원색 위주의 장식 이미지
        score -= 60.0
    elif mean_sat > 25:
        score -= 20.0
    if colorful > 0.5:         # 화면 대부분이 유채색 → 도면이 아닐 가능성
        score -= 40.0
    if area and area < 2500:   # 썸네일 기준 50x50 미만
        score -= 25.0
    return score


def select_representative_images(rows: list[dict]) -> list[str | None]:
    """각 건의 대표도면을 고른다.

    EUIPO 등록증의 표지 그래픽·EU 깃발 배경처럼 도면이 아닌 서식 이미지가
    선택되지 않도록, 후보마다 점수를 매겨 가장 도면다운 것을 고른다.
    (반복 등장 횟수 + 색상 특성. 반복 횟수만으로는 소수 건에만 들어간
    장식 이미지를 걸러내지 못하기 때문)
    """
    from collections import Counter

    per_row = [_drawing_paths(r.get("drawings_json")) for r in rows]

    counts: Counter = Counter()
    row_hashes: list[list[tuple[str, str | None]]] = []
    for paths in per_row:
        hashed = [(p, _file_hash(p)) for p in paths]
        row_hashes.append(hashed)
        counts.update({h for _, h in hashed if h})  # 건 단위로 1회만 집계

    picked: list[str | None] = []
    for hashed in row_hashes:
        if not hashed:
            picked.append(None)
            continue
        best = max(
            hashed,
            key=lambda ph: _image_score(ph[0], counts.get(ph[1], 1)),
        )
        picked.append(best[0])
    return picked


def representative_image(drawings_json: str | None) -> str | None:
    """도면 목록에서 대표도면 파일 경로를 고른다 (없으면 첫 번째 도면)."""
    if not drawings_json:
        return None
    try:
        drawings = json.loads(drawings_json)
    except (json.JSONDecodeError, TypeError):
        return None
    if not isinstance(drawings, list) or not drawings:
        return None

    def path_of(d):
        return d.get("file_path") or d.get("url") if isinstance(d, dict) else None

    for d in drawings:
        if isinstance(d, dict) and d.get("drawing_type") == "representative":
            p = path_of(d)
            if p and Path(p).exists():
                return p
    for d in drawings:
        p = path_of(d)
        if p and Path(p).exists():
            return p
    return None


def to_excel_with_images_bytes(rows: list[dict], image_paths: list[str | None],
                               thumb_px: int = 110) -> bytes:
    """분류 결과를 Excel로 만들되 '대표도면' 열에 이미지를 삽입한다."""
    import io

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
    from PIL import Image as PILImage

    wb = Workbook()
    ws = wb.active
    ws.title = "분류결과"

    # 한글 폰트를 지정하지 않으면 일부 환경에서 한글이 중국어 폰트로 대체되어
    # 깨져 보일 수 있다.
    from openpyxl.styles import Font

    korean_font = Font(name="Malgun Gothic", size=10)

    headers = ["대표도면"] + list(rows[0].keys()) if rows else ["대표도면"]
    ws.append(headers)

    for r in rows:
        ws.append([""] + [r.get(h, "") for h in headers[1:]])

    # 이미지 열 크기 지정 (열 너비 단위 ≈ px/7, 행 높이 단위 = pt ≈ px*0.75)
    ws.column_dimensions["A"].width = thumb_px / 7
    for i, h in enumerate(headers[1:], start=2):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(h)) + 6, 14), 45)

    for row in ws.iter_rows():
        for cell in row:
            cell.font = korean_font

    buffers = []  # 저장 전까지 스트림을 살려둬야 함
    for idx, src in enumerate(image_paths):
        row_no = idx + 2
        if not src:
            continue
        try:
            with PILImage.open(src) as im:
                im = im.convert("RGB")
                im.thumbnail((thumb_px, thumb_px))
                buf = io.BytesIO()
                im.save(buf, format="PNG")
            buf.seek(0)
            buffers.append(buf)
            img = XLImage(buf)
            ws.add_image(img, f"A{row_no}")
            ws.row_dimensions[row_no].height = thumb_px * 0.75
        except Exception:
            continue  # 이미지 하나가 깨져도 엑셀 생성은 계속

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_excel_bytes(rows: list[dict]) -> bytes:
    """dict 목록을 Excel 파일 바이트로 변환 (다운로드용)."""
    import io

    import pandas as pd

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, index=False, sheet_name="분류결과")
    return buf.getvalue()
