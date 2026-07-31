"""분류 결과 내보내기(Excel·대표도면 삽입) 테스트."""
import io
import json
import zipfile

import pytest
from openpyxl import load_workbook
from PIL import Image

from src.utils.export import (
    representative_image,
    to_excel_bytes,
    to_excel_with_images_bytes,
)


@pytest.fixture
def images(tmp_path):
    paths = []
    for i, color in enumerate(["red", "green", "blue"]):
        p = tmp_path / f"draw{i}.png"
        Image.new("RGB", (900, 700), color).save(p)
        paths.append(str(p))
    return paths


def test_representative_drawing_preferred(images):
    drawings_json = json.dumps([
        {"file_path": images[1], "drawing_type": "perspective"},
        {"file_path": images[0], "drawing_type": "representative"},
    ])
    assert representative_image(drawings_json) == images[0]


def test_falls_back_to_first_existing_file(images):
    drawings_json = json.dumps([
        {"file_path": "/없는/경로.png", "drawing_type": "front"},
        {"file_path": images[2], "drawing_type": "front"},
    ])
    assert representative_image(drawings_json) == images[2]


def test_representative_image_handles_empty_and_invalid():
    assert representative_image(None) is None
    assert representative_image("[]") is None
    assert representative_image("깨진 JSON") is None


def test_excel_embeds_images_and_survives_broken_paths(images):
    rows = [{"출원번호": f"A{i}", "물품명": "Transfer robot"} for i in range(3)]
    data = to_excel_with_images_bytes(rows, [images[0], None, "/없는/파일.png"])

    assert data[:2] == b"PK"  # 유효한 xlsx
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        media = [n for n in z.namelist() if n.startswith("xl/media/")]
    assert len(media) == 1  # 유효한 이미지 1건만 삽입, 나머지는 건너뜀

    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["대표도면", "출원번호", "물품명"]
    assert ws["B2"].value == "A0"
    assert ws.row_dimensions[2].height > 0  # 이미지가 보이도록 행 높이 확대


def test_excel_without_images_still_works():
    rows = [{"출원번호": "A0", "주분류": "곡선형"}]
    data = to_excel_bytes(rows)
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["출원번호", "주분류"]


def test_boilerplate_cover_image_excluded(tmp_path):
    """모든 건에 반복되는 표지·로고 이미지는 대표도면에서 제외한다."""
    import json
    import shutil
    from src.utils.export import select_representative_images

    cover = tmp_path / "cover.png"
    Image.new("RGB", (300, 200), (20, 60, 200)).save(cover)

    rows = []
    for i in range(10):
        d = tmp_path / f"p{i}"
        d.mkdir()
        c = d / "img0.png"
        shutil.copy(cover, c)            # 전 건 공통(표지)
        uniq = d / "img1.png"            # 건별 고유 도면
        Image.new("RGB", (400, 300), (i * 20 % 255, 128, 64)).save(uniq)
        rows.append({"drawings_json": json.dumps([
            {"file_path": str(c), "drawing_type": "representative"},
            {"file_path": str(uniq), "drawing_type": "front"},
        ])})

    picked = select_representative_images(rows)
    assert all(p and "img1" in p for p in picked), picked


def test_falls_back_when_only_boilerplate_available(tmp_path):
    """도면이 공통 이미지뿐이면 빈칸 대신 그것이라도 사용."""
    import json
    from src.utils.export import select_representative_images

    cover = tmp_path / "cover.png"
    Image.new("RGB", (300, 200), "blue").save(cover)
    rows = [{"drawings_json": json.dumps(
        [{"file_path": str(cover), "drawing_type": "representative"}])} for _ in range(4)]
    assert select_representative_images(rows)[0] is not None


def test_select_handles_missing_and_invalid(tmp_path):
    from src.utils.export import select_representative_images
    rows = [{"drawings_json": None}, {"drawings_json": "깨진 JSON"},
            {"drawings_json": '[{"file_path": "/없는/파일.png"}]'}]
    assert select_representative_images(rows) == [None, None, None]


def _make_flag(path):
    """EU 깃발처럼 채도가 높은 장식 이미지."""
    im = Image.new("RGB", (200, 140))
    im.putdata([(20, 60, 220) if (x // 10 + y // 10) % 2 else (250, 220, 20)
                for y in range(140) for x in range(200)])
    im.save(path)


def _make_drawing(path, seed=0):
    """도면처럼 흰 바탕에 어두운 선만 있는 이미지."""
    im = Image.new("RGB", (300, 220), "white")
    px = im.load()
    for x in range(20, 280):
        px[x, 60 + (seed % 30)] = (30, 30, 30)
        px[x, 160] = (30, 30, 30)
    im.save(path)


def test_decorative_image_rejected_even_when_rarely_repeated(tmp_path):
    """소수 건에만 들어간 장식 이미지(EU 깃발)도 대표도면에서 제외한다.

    반복 횟수만 보던 방식은 3/30건(10%)짜리 장식을 놓쳤으므로,
    색상 특성까지 함께 평가해야 한다.
    """
    import json
    import shutil
    from src.utils.export import select_representative_images

    flag = tmp_path / "flag.png"
    _make_flag(flag)

    rows = []
    for i in range(30):
        d = tmp_path / f"p{i}"
        d.mkdir()
        imgs = []
        if i in (0, 2, 3):                       # 3건에만 등장하는 장식
            f = d / "a_flag.png"
            shutil.copy(flag, f)
            imgs.append(str(f))
        real = d / "b_real.png"                  # 건별 고유 도면
        _make_drawing(real, seed=i)
        imgs.append(str(real))
        rows.append({"drawings_json": json.dumps([
            {"file_path": p, "drawing_type": "representative" if j == 0 else "front"}
            for j, p in enumerate(imgs)
        ])})

    picked = select_representative_images(rows)
    assert all(p and "b_real" in p for p in picked), [picked[i] for i in (0, 2, 3)]


def test_drawing_scores_higher_than_decoration(tmp_path):
    from src.utils.export import _image_score

    flag = tmp_path / "flag.png"
    _make_flag(flag)
    drawing = tmp_path / "drawing.png"
    _make_drawing(drawing)

    # 장식은 반복이 적어도(2건) 도면(고유 1건)보다 낮은 점수여야 함
    assert _image_score(str(flag), 2) < _image_score(str(drawing), 1)
